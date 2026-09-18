"""Deterministic document layouts: style supplied content without rewriting facts."""
from __future__ import annotations

from dataclasses import dataclass
import io
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from .fonts import has_devanagari
from .quality import note

PRESETS = {
    "general": dict(margin=1.0, spacing=1.15, after=7, heading=15, title=23),
    "formal": dict(margin=1.0, spacing=1.15, after=9, heading=14, title=18),
    "report": dict(margin=0.85, spacing=1.2, after=7, heading=16, title=25),
    "resume": dict(margin=0.65, spacing=1.05, after=4, heading=12, title=21),
    "custom": dict(margin=1.0, spacing=1.15, after=7, heading=15, title=23),
}
KNOWN_HEADINGS = {
    "summary", "professional summary", "profile", "objective", "experience",
    "work experience", "professional experience", "employment history", "education",
    "skills", "technical skills", "core competencies", "projects", "certifications",
    "achievements", "awards", "languages", "volunteering", "publications", "interests",
    "executive summary", "introduction", "background", "objectives", "scope",
    "methodology", "findings", "results", "discussion", "recommendations",
    "conclusion", "next steps", "references", "appendix",
}


def heading_key(text: str) -> str:
    return " ".join(text.strip().rstrip(":").split()).casefold()


@dataclass
class Block:
    text: str
    kind: str = "body"
    level: int = 1


def parse_blocks(text: str, section_order: list[str]) -> list[Block]:
    """Only explicit syntax and exact heading labels have structural meaning."""
    headings = KNOWN_HEADINGS | {heading_key(s) for s in section_order}
    blocks = []
    for line in text.replace("\r\n", "\n").replace("\r", "\n").split("\n"):
        stripped = line.strip()
        heading = re.match(r"^(#{1,3})\s+(.+)$", stripped)
        bullet = re.match(r"^[-*+•]\s+(.+)$", stripped)
        if heading:
            blocks.append(Block(heading[2], "heading", len(heading[1])))
        elif stripped and heading_key(stripped) in headings:
            blocks.append(Block(stripped, "heading"))
        elif bullet:
            blocks.append(Block(bullet[1], "bullet"))
        elif re.match(r"^\d+[.)]\s+", stripped):
            blocks.append(Block(line, "numbered"))
        else:
            blocks.append(Block(line, "body" if stripped else "blank"))
    return blocks


def order_sections(blocks: list[Block], order: list[str]) -> list[Block]:
    if not order:
        return blocks
    keys = [heading_key(s) for s in order]
    if any(not key for key in keys) or len(keys) != len(set(keys)):
        raise ValueError("Custom section names must be non-empty and unique.")
    preamble, groups = [], []
    for block in blocks:
        if block.kind == "heading" and block.level == 1:
            groups.append([block])
        elif groups:
            groups[-1].append(block)
        else:
            preamble.append(block)
    existing = {heading_key(group[0].text) for group in groups}
    missing = [name for name in order if heading_key(name) not in existing]
    if missing:
        raise ValueError("Custom sections not found in your content: " + ", ".join(missing) +
                         ". Add each as a standalone heading or '# Heading'. No content was generated.")
    ordered = [group for key in keys for group in groups if heading_key(group[0].text) == key]
    remaining = [group for group in groups if heading_key(group[0].text) not in keys]
    if remaining:
        note("Sections not listed in your custom order were kept at the end; no supplied sections were discarded.")
    return preamble + [block for group in ordered + remaining for block in group]


def build_document(text: str, font: str, font_size: float, template: str = "general",
                   title: str = "", custom_format: dict | None = None) -> bytes:
    if template not in PRESETS:
        raise ValueError("Choose general, formal, report, resume or custom format.")
    # Word XML cannot represent these control characters. Reject rather than dropping content.
    if re.search(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", text + title):
        raise ValueError("Your text contains unsupported control characters. Remove them and try again.")
    custom = custom_format or {}
    options = PRESETS[template].copy()
    section_order = custom.get("section_order", []) if template == "custom" else []
    page_size, alignment = "A4", "left"
    if template == "custom":
        options.update(margin=custom.get("margin_inches", 1.0),
                       spacing=custom.get("line_spacing", 1.15),
                       heading=custom.get("heading_size", 15))
        page_size = custom.get("page_size", "A4")
        alignment = custom.get("heading_alignment", "left")
    blocks = order_sections(parse_blocks(text, section_order), section_order)
    doc = Document()
    doc.core_properties.title = title.strip()
    doc.core_properties.subject = f"{template.title()} document"
    section = doc.sections[0]
    section.page_width = Inches(8.5) if page_size == "Letter" else Pt(595.276)
    section.page_height = Inches(11) if page_size == "Letter" else Pt(841.89)
    for edge in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, edge, Inches(options["margin"]))

    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "List Bullet"):
        style = doc.styles[name]
        style.font.name = font
        style.font.size = Pt(font_size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        props = style.element.get_or_add_rPr()
        families = props.find(qn("w:rFonts"))
        if families is None:
            from docx.oxml import OxmlElement
            families = OxmlElement("w:rFonts"); props.insert(0, families)
        for attr in ("asciiTheme", "hAnsiTheme", "cstheme", "eastAsiaTheme"):
            families.attrib.pop(qn(f"w:{attr}"), None)
        families.set(qn("w:ascii"), font)
        families.set(qn("w:hAnsi"), font)
        families.set(qn("w:cs"), font)
        style.paragraph_format.line_spacing = options["spacing"]
        style.paragraph_format.space_after = Pt(options["after"])
        style.paragraph_format.widow_control = True
    for level in (1, 2, 3):
        style = doc.styles[f"Heading {level}"]
        style.font.size = Pt(max(font_size, options["heading"] - (level - 1)))
        style.font.bold = True
        style.paragraph_format.space_before = Pt(12 if template != "resume" else 8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if alignment == "center" else WD_ALIGN_PARAGRAPH.LEFT
    doc.styles["Title"].font.size = Pt(max(options["title"], font_size + 4))
    doc.styles["Title"].font.bold = True
    doc.styles["Title"].paragraph_format.keep_with_next = True
    doc.styles["Title"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER if alignment == "center" else WD_ALIGN_PARAGRAPH.LEFT

    def add_text(paragraph, value):
        # Keep script-specific fonts local to their characters, including mixed-language lines.
        parts = re.split(r"([\u0900-\u097f\u1cd0-\u1cff\ua8e0-\ua8ff]+)", value)
        for part in parts:
            if not part and value:
                continue
            run = paragraph.add_run(part)
            run.font.name = font
            families = run._element.get_or_add_rPr().find(qn("w:rFonts"))
            if has_devanagari(part):
                families.set(qn("w:cs"), "Noto Sans Devanagari")

    if title.strip():
        add_text(doc.add_paragraph(style="Title"), title.strip())
    for block in blocks:
        style = f"Heading {block.level}" if block.kind == "heading" else "List Bullet" if block.kind == "bullet" else "Normal"
        paragraph = doc.add_paragraph(style=style)
        add_text(paragraph, block.text)
        if block.kind in ("bullet", "numbered"):
            paragraph.paragraph_format.left_indent = Inches(0.18)
            paragraph.paragraph_format.first_line_indent = Inches(-0.18)
        if block.kind == "blank":
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1
            paragraph.runs[0].font.size = Pt(4)
    if template == "resume":
        note("Resume uses one column, standard headings and body text without tables or text boxes. ATS results also depend on the job requirements and your content; no score is guaranteed.")
        if font_size < 10 or font_size > 12:
            note("Resume body text is usually easier to read at 10–12 pt. Your selected font size was retained.")
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()

"""
Image/Scan → DOCX

Strategy:
  1. Load image, run OCR (Tesseract, eng+hin).
  2. Group OCR lines into paragraphs by vertical gap.
  3. Detect heading lines by font-size heuristic from bounding box height.
  4. Embed original image at top of doc for visual reference.
  5. Append reconstructed text below.
"""
from __future__ import annotations

import io

from docx import Document
from docx.shared import Inches, Pt, RGBColor

from .shared import (
    ConversionError, load_image, image_to_jpeg_bytes,
    ocr_image, validate_docx,
)


def convert(data: bytes) -> bytes:
    im = load_image(data)
    jpeg = image_to_jpeg_bytes(im)

    ocr_lines = ocr_image(im)
    if not ocr_lines:
        raise ConversionError("OCR found no text in this image")

    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

    # Original image at top
    try:
        doc.add_picture(io.BytesIO(jpeg), width=Inches(5.5))
    except Exception:
        pass
    doc.add_paragraph()

    # Compute median line height for heading heuristic
    heights = [ln["h"] for ln in ocr_lines]
    med_h = sorted(heights)[len(heights) // 2] if heights else 20

    groups = _group_lines(ocr_lines)
    for group in groups:
        text = " ".join(ln["text"] for ln in group)
        max_h = max(ln["h"] for ln in group)
        is_bold = text.isupper() or max_h > med_h * 1.5

        if max_h >= med_h * 1.8:
            doc.add_heading(text, level=1)
        elif max_h >= med_h * 1.35 or is_bold:
            doc.add_heading(text, level=2)
        else:
            para = doc.add_paragraph()
            run = para.add_run(text)
            run.font.size = Pt(max(8, min(36, max_h * 0.8)))
            run.bold = is_bold

    buf = io.BytesIO()
    doc.save(buf)
    result = buf.getvalue()
    validate_docx(result)
    return result


def _group_lines(lines: list[dict]) -> list[list[dict]]:
    """Group OCR lines into paragraph groups by y-gap."""
    if not lines:
        return []
    lines = sorted(lines, key=lambda l: (l["y"], l["x"]))
    groups: list[list[dict]] = [[lines[0]]]
    for ln in lines[1:]:
        last = groups[-1][-1]
        gap = ln["y"] - (last["y"] + last["h"])
        same_line = abs(ln["y"] - last["y"]) < last["h"] * 0.5
        if same_line or gap < max(last["h"] * 0.8, 6):
            groups[-1].append(ln)
        else:
            groups.append([ln])
    return groups

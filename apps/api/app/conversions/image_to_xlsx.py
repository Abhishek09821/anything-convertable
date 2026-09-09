"""
Image/Scan → XLSX

Strategy:
  1. OCR the image.
  2. Cluster x-positions into columns (20px tolerance).
  3. Group by y-row (same line = same row).
  4. Write to worksheet with basic formatting.
  5. Coerce numeric strings to numbers.
"""
from __future__ import annotations

import io
import re

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .shared import ConversionError, load_image, ocr_image, validate_xlsx


def convert(data: bytes) -> bytes:
    im = load_image(data)
    lines = ocr_image(im)
    if not lines:
        raise ConversionError("OCR found no text in this image")

    grid = _lines_to_grid(lines)
    if not grid:
        raise ConversionError("Could not detect tabular structure")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"

    hdr_fill = PatternFill("solid", fgColor="1F4E79")
    hdr_font = Font(bold=True, color="FFFFFF", size=10)
    alt_fill = PatternFill("solid", fgColor="EBF3FB")
    body_font = Font(size=9)

    for ri, row in enumerate(grid):
        for ci, val in enumerate(row):
            cell = ws.cell(row=ri + 1, column=ci + 1)
            cell.value = _coerce(val)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if ri == 0:
                cell.font = hdr_font
                cell.fill = hdr_fill
            elif ri % 2 == 0:
                cell.fill = alt_fill
                cell.font = body_font
            else:
                cell.font = body_font

    for col in ws.columns:
        max_len = max((len(str(c.value or "")) for c in col), default=0)
        ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 4, 50)

    if ws.max_row > 1:
        ws.freeze_panes = ws["A2"]

    buf = io.BytesIO()
    wb.save(buf)
    result = buf.getvalue()
    validate_xlsx(result)
    return result


def _lines_to_grid(lines: list[dict]) -> list[list[str]]:
    """Convert OCR line dicts into a 2-D grid via column clustering."""
    if not lines:
        return []

    # Collect x origins
    xs = sorted(set(round(ln["x"], 0) for ln in lines))
    col_anchors: list[float] = [xs[0]]
    for x in xs[1:]:
        if x - col_anchors[-1] > 20:
            col_anchors.append(x)

    n_cols = len(col_anchors)

    # Group lines by row (y within tolerance)
    row_groups: dict[int, list[dict]] = {}
    for ln in lines:
        ry = round(ln["y"] / 5) * 5
        row_groups.setdefault(ry, []).append(ln)

    grid: list[list[str]] = []
    for ry in sorted(row_groups):
        row_cells = [""] * n_cols
        for ln in sorted(row_groups[ry], key=lambda l: l["x"]):
            ci = min(range(n_cols), key=lambda i: abs(col_anchors[i] - ln["x"]))
            row_cells[ci] = (row_cells[ci] + " " + ln["text"]).strip()
        grid.append(row_cells)

    return grid


def _coerce(val: str):
    v = str(val).strip().replace(",", "")
    stripped = re.sub(r"^[₹$€£¥\s]+", "", v).strip()
    if not stripped:
        return val or None
    try:
        return int(stripped)
    except ValueError:
        pass
    try:
        return float(stripped)
    except ValueError:
        pass
    return val

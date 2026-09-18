"""PDF to PowerPoint: complete rendered pages, or editable text over a preserved graphics layer."""
from __future__ import annotations

import io
import re
from typing import Any

import fitz  # PyMuPDF
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Length, Pt

from .fonts import has_devanagari, get_unicode_font_path


# ── Constants ──────────────────────────────────────────────────────────────────
EMU_PER_PT = 12700


# ── Font helpers ───────────────────────────────────────────────────────────────

def _normalize_font_choice(font: str) -> str:
    """Normalize user-requested font option."""
    f = (font or "original").strip().lower().replace("-", "_").replace(" ", "_")
    if "times" in f:
        return "Times New Roman"
    if "arial" in f:
        return "Arial"
    if "calibri" in f:
        return "Calibri"
    if "georgia" in f:
        return "Georgia"
    return "original"


def _clean_pdf_font_name(raw_name: str | None) -> str:
    """Strip PDF subset prefix like 'ABCDEF+Calibri-Bold' -> 'Calibri'."""
    if not raw_name:
        return "Arial"
    name = raw_name
    if "+" in name:
        name = name.split("+", 1)[1]
    name = re.sub(r"[-,]?(BoldItalic|BoldOblique|Bold|Italic|Oblique|Regular|Roman|PSMT|MT)$", "", name)
    name = {"TimesNewRomanPS": "Times New Roman", "ArialMT": "Arial"}.get(name, name)
    return name.strip() or "Arial"


def _get_pptx_unicode_font_name() -> str:
    """Return font family name installed on host system for Devanagari/Unicode."""
    font_path = get_unicode_font_path()
    if font_path:
        base = font_path.lower()
        if "noto" in base:
            return "Noto Sans Devanagari"
        if "arial" in base:
            return "Arial Unicode MS"
        if "sangam" in base or "devanagari" in base:
            return "Devanagari Sangam MN"
        if "mangal" in base:
            return "Mangal"
    return "Arial Unicode MS"


# ── Helpers ────────────────────────────────────────────────────────────────────

def _pt_to_emu(pt_val: float) -> int:
    return int(pt_val * EMU_PER_PT)


def _hex_to_rgb(hex_int: int) -> RGBColor:
    r = (hex_int >> 16) & 0xFF
    g = (hex_int >> 8)  & 0xFF
    b =  hex_int        & 0xFF
    return RGBColor(r, g, b)


def _add_text_box(slide: Any, text: str, x_emu: int, y_emu: int,
                  w_emu: int, h_emu: int,
                  font_size_pt: float, bold: bool, italic: bool,
                  color_int: int, align_str: str,
                  target_font_name: str,
                  unicode_font_name: str):
    """Add an editable text box at the given position."""
    if not text.strip():
        return
    if w_emu < 1 or h_emu < 1:
        return

    # Minimum size guards
    w_emu = max(w_emu, _pt_to_emu(15))
    h_emu = max(h_emu, _pt_to_emu(12))

    txBox = slide.shapes.add_textbox(x_emu, y_emu, w_emu, h_emu)
    tf = txBox.text_frame
    tf.word_wrap = False
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0

    p = tf.paragraphs[0]

    # Alignment
    al = align_str.lower() if align_str else ""
    if "center" in al:
        p.alignment = PP_ALIGN.CENTER
    elif "right" in al:
        p.alignment = PP_ALIGN.RIGHT
    else:
        p.alignment = PP_ALIGN.LEFT

    run = p.add_run()
    run.text = text
    run.font.size = Pt(max(6.0, font_size_pt))
    run.font.bold = bold
    run.font.italic = italic

    # Apply font choice at python-pptx and OOXML levels
    from pptx.oxml.ns import qn
    is_unicode = has_devanagari(text)
    latin_font = target_font_name if (target_font_name and target_font_name != "original") else "Arial"
    run.font.name = unicode_font_name if is_unicode else latin_font

    try:
        rPr = run._element.get_or_add_rPr()
        latin_el = rPr.find(qn("a:latin"))
        if latin_el is None:
            latin_el = rPr.makeelement(qn("a:latin"), {})
            rPr.append(latin_el)
        latin_el.set("typeface", latin_font)

        cs_el = rPr.find(qn("a:cs"))
        if cs_el is None:
            cs_el = rPr.makeelement(qn("a:cs"), {})
            rPr.append(cs_el)
        cs_el.set("typeface", unicode_font_name)
    except Exception:
        pass

    try:
        run.font.color.rgb = _hex_to_rgb(color_int)
    except Exception:
        pass


def _merge_line_spans(spans: list[dict]) -> list[dict]:
    """Merge horizontally adjacent spans on the same line to reduce shape count."""
    if len(spans) <= 1:
        return spans

    spans.sort(key=lambda s: s["bbox"][0])
    merged = [spans[0].copy()]

    for span in spans[1:]:
        prev = merged[-1]
        prev_x1 = prev["bbox"][2]
        curr_x0 = span["bbox"][0]
        gap = curr_x0 - prev_x1
        same_size = abs(prev.get("size", 12) - span.get("size", 12)) < 1.5
        same_flags = prev.get("flags", 0) == span.get("flags", 0)
        same_color = prev.get("color", 0) == span.get("color", 0)
        same_font = prev.get("font") == span.get("font")

        if 0 <= gap < prev.get("size", 12) * 0.75 and same_size and same_flags and same_color and same_font:
            sep = " " if gap > 1.0 else ""
            prev["text"] = prev["text"] + sep + span["text"]
            prev["bbox"] = (
                prev["bbox"][0],
                min(prev["bbox"][1], span["bbox"][1]),
                span["bbox"][2],
                max(prev["bbox"][3], span["bbox"][3]),
            )
        else:
            merged.append(span.copy())

    return merged


# ── Main ───────────────────────────────────────────────────────────────────────

def convert(data: bytes, font: str = "original", fidelity: str = "editable") -> bytes:
    from .quality import note
    if fidelity not in ("appearance", "editable"):
        raise ValueError("Choose appearance or editable conversion mode.")
    prs = Presentation()
    selected_font = _normalize_font_choice(font)
    unicode_font = _get_pptx_unicode_font_name()
    with fitz.open(stream=data, filetype="pdf") as pdf:
        if pdf.needs_pass:
            raise ValueError("This PDF is password protected. Upload an unlocked copy.")
        if not len(pdf):
            raise ValueError("The PDF contains no pages.")
        # PPTX supports one page size for the entire deck; center other sizes without stretching.
        first = pdf[0].rect
        canvas_w, canvas_h = first.width, first.height
        # OOXML slide dimensions have a 56-inch maximum.
        canvas_scale = min(1, 4032 / max(canvas_w, canvas_h))
        canvas_w *= canvas_scale
        canvas_h *= canvas_scale
        prs.slide_width = Length(_pt_to_emu(canvas_w))
        prs.slide_height = Length(_pt_to_emu(canvas_h))
        for page in pdf:
            rect = page.rect
            scale = min(canvas_w / rect.width, canvas_h / rect.height)
            left, top = (canvas_w - rect.width*scale)/2, (canvas_h - rect.height*scale)/2
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            native = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)
            # A rendered graphics layer preserves masks, clipped/rotated images, paths and stacking.
            # For editable output remove only text, retaining the complete graphics underneath.
            with fitz.open() as background:
                background.insert_pdf(pdf, from_page=page.number, to_page=page.number)
                bg = background[0]
                has_text = any(span.get("type") != 3 and span.get("chars") for span in page.get_texttrace())
                horizontal = all(tuple(line.get("dir", (1, 0))) == (1, 0)
                    for block in native.get("blocks", []) for line in block.get("lines", []))
                editable = fidelity == "editable" and has_text and page.rotation == 0 and horizontal
                if editable:
                    bg.add_redact_annot(bg.rect, fill=False)
                    bg.apply_redactions(images=0, graphics=0, text=0)
                dpi = min(300, max(72, int(72 * (24_000_000/(rect.width*rect.height))**.5)))
                pix = bg.get_pixmap(dpi=dpi, alpha=False, colorspace=fitz.csRGB)
                slide.shapes.add_picture(io.BytesIO(pix.tobytes("png")),
                    Length(_pt_to_emu(left)), Length(_pt_to_emu(top)),
                    Length(_pt_to_emu(rect.width*scale)), Length(_pt_to_emu(rect.height*scale)))
            if not editable:
                if fidelity == "editable":
                    note("Scanned, rotated or angled-text pages were preserved as images to avoid layout loss; their text is not editable.")
                continue
            for block in native.get("blocks", []):
                for line in block.get("lines", []):
                    for span in _merge_line_spans(line.get("spans", [])):
                        x0, y0, x1, y1 = span["bbox"]
                        flags = span.get("flags", 0)
                        family = selected_font if selected_font != "original" else _clean_pdf_font_name(span.get("font"))
                        _add_text_box(slide, span.get("text", ""),
                            _pt_to_emu(left+x0*scale), _pt_to_emu(top+y0*scale),
                            _pt_to_emu((x1-x0)*scale*1.03), _pt_to_emu((y1-y0)*scale),
                            span.get("size", 12)*scale, bool(flags & 16), bool(flags & 2),
                            span.get("color", 0), "left", family, unicode_font)
            if page.get_text().strip():
                note("Text is editable; graphics are preserved as a single image layer. Font substitution and complex text placement may differ.")
        if fidelity == "appearance":
            note("Pages are preserved as high-resolution slide images. Text and individual graphics are not editable.")
        if any(abs(p.rect.width-first.width) > 1 or abs(p.rect.height-first.height) > 1 for p in pdf):
            note("Mixed page sizes were fitted to a single slide size without stretching.")
    output = io.BytesIO()
    prs.save(output)
    return output.getvalue()

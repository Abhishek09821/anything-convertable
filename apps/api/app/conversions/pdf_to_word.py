"""PDF to Word: native layout reconstruction plus per-page OCR and source scan preservation."""
from __future__ import annotations

import io
import os
import tempfile

import fitz  # PyMuPDF
from docx import Document
from docx.shared import Pt
from pdf2docx import Converter


# ── Font options & normalization ──────────────────────────────────────────────

def _normalize_font_choice(font: str) -> str:
    f = (font or "original").strip().lower().replace("-", "_").replace(" ", "_")
    if "times" in f:
        return "Times New Roman"
    if "arial" in f:
        return "Arial"
    if "calibri" in f:
        return "Calibri"
    if "georgia" in f:
        return "Georgia"
    return "original"


def _apply_font_to_docx(docx_bytes: bytes, font_choice: str) -> bytes:
    from .office import apply_office_font
    target = _normalize_font_choice(font_choice)
    return apply_office_font(docx_bytes, target)


# ── pdf2docx converter ────────────────────────────────────────────────────────

def _convert_with_pdf2docx(data: bytes) -> bytes:
    """Convert searchable PDF to DOCX using tuned pdf2docx layout reconstruction."""
    with tempfile.NamedTemporaryFile(suffix=".pdf", delete=False) as f_in:
        f_in.write(data)
        pdf_path = f_in.name
    with tempfile.NamedTemporaryFile(suffix=".docx", delete=False) as f_out:
        docx_path = f_out.name
    try:
        cv = Converter(pdf_path)
        # Optimal parameters for layout fidelity: preserve tables, paragraphs, and high-res images
        try:
            cv.convert(
                docx_path, start=0,
                parse_lattice_table=True, parse_stream_table=True,
                clip_image_res_ratio=4.0, ignore_page_error=False,
                multi_processing=False,
            )
        finally:
            cv.close()
        with open(docx_path, "rb") as f:
            return f.read()
    finally:
        try:
            os.unlink(pdf_path)
        except Exception:
            pass
        try:
            os.unlink(docx_path)
        except Exception:
            pass


def _validate_docx(data: bytes) -> bool:
    """Verify the output is a valid DOCX (ZIP with correct structure)."""
    import zipfile
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            return "[Content_Types].xml" in z.namelist()
    except Exception:
        return False


# ── Main ───────────────────────────────────────────────────────────────────────

def _page_document(page, font: str) -> bytes:
    """Keep a source scan alongside OCR text so graphics are never discarded."""
    from .ocr import scan_lines
    from .quality import note
    doc = Document()
    section = doc.sections[0]
    section.page_width = Pt(page.rect.width)
    section.page_height = Pt(page.rect.height)
    for edge in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(section, edge, Pt(18))
    lines = scan_lines(page)
    for line in lines:
        run = doc.add_paragraph(line["text"]).runs[0]
        run.font.size = Pt(max(7, min(36, (line["bbox"][3] - line["bbox"][1]) * 1.2)))
    if lines:
        doc.add_page_break()
    # Preserve diagrams, photographs, stamps, signatures and the original typography.
    pix = page.get_pixmap(dpi=200, colorspace=fitz.csRGB, alpha=False)
    width = min(page.rect.width - 36, (page.rect.height - 48) * page.rect.width / page.rect.height)
    doc.add_picture(io.BytesIO(pix.tobytes("png")), width=Pt(width))
    note("Scanned pages include recognized editable text followed by a source-page image. OCR layout and font sizes are approximate.")
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def convert(data: bytes, font: str = "original") -> bytes:
    """Analyze each page; never let a native page suppress OCR on other pages."""
    from docxcompose.composer import Composer
    from .quality import note
    with fitz.open(stream=data, filetype="pdf") as pdf:
        if pdf.needs_pass:
            raise ValueError("This PDF is password protected. Upload an unlocked copy.")
        if not len(pdf):
            raise ValueError("The PDF contains no pages.")
        native = [bool(page.get_text().strip()) and any(
            span.get("type") != 3 and span.get("chars")
            for span in page.get_texttrace()) for page in pdf]
        if all(native):
            result = _convert_with_pdf2docx(data)
        else:
            composer = None
            for index, page in enumerate(pdf):
                if native[index]:
                    with fitz.open() as single:
                        single.insert_pdf(pdf, from_page=index, to_page=index)
                        part = _convert_with_pdf2docx(single.tobytes())
                else:
                    part = _page_document(page, font)
                document = Document(io.BytesIO(part))
                if composer is None:
                    composer = Composer(document)
                else:
                    composer.doc.add_page_break()
                    composer.append(document)
            out = io.BytesIO()
            composer.save(out)
            result = out.getvalue()
    if font != "original":
        note("Changing the font can change line breaks and pagination.")
    note("Editable Word layout is reconstructed. Fonts must be available on the device opening the document.")
    result = _apply_font_to_docx(result, font)
    if not _validate_docx(result):
        raise ValueError("The conversion did not produce a valid Word document.")
    return result

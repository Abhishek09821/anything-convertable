"""
shared.py — extraction primitives reused by all 10 converters.

Provides:
  - extract_pdf_pages()    rich page data from a PDF (native text + images)
  - ocr_image()            run Tesseract on a PIL Image, return line-level results
  - classify_document()    detect invoice / resume / screenshot / generic
  - ConversionError        typed exception
  - ValidationError        typed exception for output validation failures
  - validate_docx()        open + check text non-empty
  - validate_xlsx()        open + check cells
  - validate_pptx()        open + check slides
  - validate_pdf()         check magic bytes + size
  - validate_html()        check non-empty markup
  - validate_json()        check schema fields present
"""
from __future__ import annotations

import base64
import io
import re
from dataclasses import dataclass, field
from typing import Any

import fitz  # PyMuPDF
from PIL import Image


# ─────────────────────────────────────────────────────────────────────────────
# Custom exceptions
# ─────────────────────────────────────────────────────────────────────────────

class ConversionError(RuntimeError):
    """Raised when a converter cannot complete."""


class ValidationError(RuntimeError):
    """Raised when output validation fails."""


# ─────────────────────────────────────────────────────────────────────────────
# Data structures
# ─────────────────────────────────────────────────────────────────────────────

@dataclass
class TextSpan:
    text: str
    x: float
    y: float
    w: float
    h: float
    size: float = 12.0
    bold: bool = False
    italic: bool = False
    color: str = "#000000"


@dataclass
class TableCell:
    row: int
    col: int
    text: str
    row_span: int = 1
    col_span: int = 1


@dataclass
class TableBlock:
    rows: int
    cols: int
    cells: list[TableCell] = field(default_factory=list)
    x: float = 0
    y: float = 0
    w: float = 0
    h: float = 0

    def cell(self, r: int, c: int) -> str:
        for cell in self.cells:
            if cell.row == r and cell.col == c:
                return cell.text
        return ""

    def as_rows(self) -> list[list[str]]:
        return [[self.cell(r, c) for c in range(self.cols)] for r in range(self.rows)]


@dataclass
class PageData:
    page_num: int
    width: float
    height: float
    spans: list[TextSpan] = field(default_factory=list)
    tables: list[TableBlock] = field(default_factory=list)
    images: list[bytes] = field(default_factory=list)  # raw image bytes
    is_scanned: bool = False  # True when no native text was extracted

    def all_text(self) -> str:
        return " ".join(s.text for s in self.spans)


# ─────────────────────────────────────────────────────────────────────────────
# PDF extraction
# ─────────────────────────────────────────────────────────────────────────────

def _rgb_to_hex(rgb: int) -> str:
    r = (rgb >> 16) & 0xFF
    g = (rgb >> 8) & 0xFF
    b = rgb & 0xFF
    return f"#{r:02x}{g:02x}{b:02x}"


def extract_pdf_pages(data: bytes, ocr_scanned: bool = True) -> list[PageData]:
    """
    Extract rich page data from PDF bytes.
    For text PDFs: native text + basic table detection.
    For scanned PDFs: render to image then OCR.
    """
    doc = fitz.open(stream=data, filetype="pdf")
    pages: list[PageData] = []

    for page_num, page in enumerate(doc):
        rect = page.rect
        pd = PageData(
            page_num=page_num,
            width=rect.width,
            height=rect.height,
        )

        # ── native text extraction ──────────────────────────────────────────
        blocks = page.get_text("dict", flags=fitz.TEXT_PRESERVE_WHITESPACE)["blocks"]
        spans_found = 0
        for b in blocks:
            if b.get("type") != 0:  # 0 = text block
                continue
            for line in b.get("lines", []):
                for span in line.get("spans", []):
                    raw = span.get("text", "").strip()
                    if not raw:
                        continue
                    bbox = span["bbox"]
                    flags = span.get("flags", 0)
                    pd.spans.append(TextSpan(
                        text=raw,
                        x=bbox[0], y=bbox[1],
                        w=bbox[2] - bbox[0],
                        h=bbox[3] - bbox[1],
                        size=round(span.get("size", 12), 1),
                        bold=bool(flags & 2**4),
                        italic=bool(flags & 2**1),
                        color=_rgb_to_hex(span.get("color", 0)),
                    ))
                    spans_found += 1

        if spans_found == 0:
            pd.is_scanned = True
            # Render page to image and OCR
            if ocr_scanned:
                mat = fitz.Matrix(2.0, 2.0)  # 2× for better OCR
                pix = page.get_pixmap(matrix=mat, colorspace=fitz.csRGB)
                img = Image.frombytes("RGB", [pix.width, pix.height], pix.samples)
                ocr_lines = ocr_image(img)
                # Scale coordinates back to PDF units
                sx = rect.width / pix.width
                sy = rect.height / pix.height
                for ln in ocr_lines:
                    pd.spans.append(TextSpan(
                        text=ln["text"],
                        x=ln["x"] * sx,
                        y=ln["y"] * sy,
                        w=ln["w"] * sx,
                        h=ln["h"] * sy,
                        size=round(max(8, ln["h"] * sy * 0.8), 1),
                        bold=ln["text"].isupper(),
                    ))
        else:
            # ── basic table detection from layout ──────────────────────────
            tables = _detect_tables_from_page(page, pd.spans)
            pd.tables.extend(tables)

        # ── embedded images ─────────────────────────────────────────────────
        for img_info in page.get_images(full=True):
            xref = img_info[0]
            try:
                base_img = doc.extract_image(xref)
                pd.images.append(base_img["image"])
            except Exception:
                pass

        # sort spans reading-order (top-down, left-right)
        pd.spans.sort(key=lambda s: (round(s.y, 1), s.x))
        pages.append(pd)

    return pages


def _detect_tables_from_page(page: Any, spans: list[TextSpan]) -> list[TableBlock]:
    """
    Use PyMuPDF's built-in table finder (available in PyMuPDF ≥ 1.23).
    Falls back to empty list if not available.
    """
    try:
        finder = page.find_tables()
        tables = []
        for t in finder.tables:
            rows = t.extract()
            if not rows or not rows[0]:
                continue
            n_rows = len(rows)
            n_cols = max(len(r) for r in rows)
            tb = TableBlock(rows=n_rows, cols=n_cols, x=t.bbox[0], y=t.bbox[1],
                            w=t.bbox[2] - t.bbox[0], h=t.bbox[3] - t.bbox[1])
            for ri, row in enumerate(rows):
                for ci, cell_text in enumerate(row):
                    tb.cells.append(TableCell(row=ri, col=ci, text=(cell_text or "").strip()))
            tables.append(tb)
        return tables
    except Exception:
        return []


# ─────────────────────────────────────────────────────────────────────────────
# OCR (Tesseract)
# ─────────────────────────────────────────────────────────────────────────────

def ocr_image(im: Image.Image, lang: str = "eng+hin") -> list[dict]:
    """
    Run Tesseract on a PIL image.
    Returns list of {text, x, y, w, h, conf} dicts grouped by line.
    lang='eng+hin' handles Hindi+English by default.
    """
    import pytesseract
    from collections import defaultdict

    # Try eng+hin; if Tesseract lacks Hindi data, fall back to eng-only
    for attempt_lang in [lang, "eng"]:
        try:
            tsv = pytesseract.image_to_data(
                im,
                lang=attempt_lang,
                output_type=pytesseract.Output.DICT,
                config="--psm 6",
            )
            break
        except Exception as exc:
            if attempt_lang == "eng":
                raise ConversionError(f"OCR failed: {exc}") from exc
            continue

    n = len(tsv["text"])
    lines: dict[tuple, list[int]] = defaultdict(list)
    for i in range(n):
        if str(tsv["text"][i]).strip():
            key = (tsv["block_num"][i], tsv["par_num"][i], tsv["line_num"][i])
            lines[key].append(i)

    results = []
    for _key, idxs in sorted(lines.items()):
        words = [str(tsv["text"][i]).strip() for i in idxs if str(tsv["text"][i]).strip()]
        if not words:
            continue
        xs = [tsv["left"][i] for i in idxs]
        ys = [tsv["top"][i] for i in idxs]
        x2s = [tsv["left"][i] + tsv["width"][i] for i in idxs]
        y2s = [tsv["top"][i] + tsv["height"][i] for i in idxs]
        confs = []
        for i in idxs:
            try:
                confs.append(max(0.0, float(tsv["conf"][i]) / 100))
            except Exception:
                confs.append(0.5)
        results.append({
            "text": " ".join(words),
            "x": min(xs), "y": min(ys),
            "w": max(x2s) - min(xs),
            "h": max(y2s) - min(ys),
            "conf": sum(confs) / len(confs) if confs else 0.5,
        })

    return results


def ocr_image_bytes(data: bytes, lang: str = "eng+hin") -> list[dict]:
    """Convenience wrapper — accepts raw image bytes."""
    im = Image.open(io.BytesIO(data)).convert("RGB")
    return ocr_image(im, lang=lang)


# ─────────────────────────────────────────────────────────────────────────────
# Document classification
# ─────────────────────────────────────────────────────────────────────────────

def classify_document(filename: str, text: str) -> str:
    """
    Returns one of: invoice, resume, screenshot, pdf, image, generic.
    """
    s = (filename + " " + text[:6000]).lower()
    if any(k in s for k in ["invoice", "bill to", "gstin", "tax invoice", "subtotal", "gst no"]):
        return "invoice"
    if any(k in s for k in ["resume", "curriculum vitae", "cv", "objective", "work experience",
                              "education", "skills", "references"]):
        return "resume"
    if filename.lower().endswith(".pdf"):
        return "pdf"
    if any(filename.lower().endswith(ext) for ext in [".png", ".jpg", ".jpeg", ".webp"]):
        # screenshots tend to have many short lines and UI text
        if len(text.split("\n")) > 20 and len(set(text.split())) > 40:
            return "screenshot"
        return "image"
    return "generic"


# ─────────────────────────────────────────────────────────────────────────────
# Image helpers
# ─────────────────────────────────────────────────────────────────────────────

def load_image(data: bytes) -> Image.Image:
    return Image.open(io.BytesIO(data)).convert("RGB")


def image_to_jpeg_bytes(im: Image.Image, quality: int = 92) -> bytes:
    buf = io.BytesIO()
    im.save(buf, format="JPEG", quality=quality)
    return buf.getvalue()


def bytes_to_data_url(raw: bytes, mime: str = "image/jpeg") -> str:
    return f"data:{mime};base64,{base64.b64encode(raw).decode()}"


# ─────────────────────────────────────────────────────────────────────────────
# Output validators
# ─────────────────────────────────────────────────────────────────────────────

def validate_docx(data: bytes, min_chars: int = 5) -> None:
    from docx import Document
    doc = Document(io.BytesIO(data))
    # Collect text from paragraphs AND table cells
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    text = " ".join(parts)
    if len(text.strip()) < min_chars:
        raise ValidationError(f"DOCX appears empty (text={repr(text[:60])})")


def validate_xlsx(data: bytes) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(data))
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                if cell.value not in (None, ""):
                    return  # found at least one non-empty cell
    raise ValidationError("XLSX has no cell values")


def validate_pptx(data: bytes) -> None:
    from pptx import Presentation
    prs = Presentation(io.BytesIO(data))
    if len(prs.slides) == 0:
        raise ValidationError("PPTX has no slides")
    found_text = False
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    if para.text.strip():
                        found_text = True
                        break
    if not found_text:
        raise ValidationError("PPTX slides contain no text")


def validate_pdf(data: bytes) -> None:
    if not data.startswith(b"%PDF"):
        raise ValidationError("Output is not a valid PDF (missing %PDF header)")
    if len(data) < 500:
        raise ValidationError(f"PDF too small ({len(data)} bytes)")


def validate_html(data: bytes | str) -> None:
    text = data.decode() if isinstance(data, bytes) else data
    if len(text.strip()) < 30:
        raise ValidationError("HTML output is empty or too small")
    if "<html" not in text.lower() and "<body" not in text.lower():
        raise ValidationError("HTML output missing <html>/<body> tags")


def validate_json(data: dict, required_keys: list[str]) -> None:
    missing = [k for k in required_keys if k not in data]
    if missing:
        raise ValidationError(f"JSON missing required keys: {missing}")


def validate_csv(data: bytes | str) -> None:
    text = data.decode() if isinstance(data, bytes) else data
    lines = [l for l in text.splitlines() if l.strip()]
    if len(lines) < 2:
        raise ValidationError(f"CSV too short ({len(lines)} lines)")

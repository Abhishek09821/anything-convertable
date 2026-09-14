"""
text_converter.py — Convert typed text / plaintext to Word (.docx) and PDF (.pdf).

Features:
- Full typography control: Times New Roman, Arial, Calibri, Georgia.
- Full Hindi / Devanagari Unicode support with complex-script Word XML tags and TrueType PDF metrics.
- Configurable font size and standard academic/business margins.
- Searchable PDF (vector text) vs Non-Searchable PDF (300 DPI high-resolution flattened raster image).
"""
from __future__ import annotations

import html
import io
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .fonts import has_devanagari, has_non_latin, register_reportlab_font
from .pdf_utils import make_non_searchable_pdf


def _normalize_font_name(font: str | None) -> str:
    f = (font or "Calibri").strip().lower().replace("-", "_").replace(" ", "_")
    if "times" in f:
        return "Times New Roman"
    if "arial" in f:
        return "Arial"
    if "calibri" in f:
        return "Calibri"
    if "georgia" in f:
        return "Georgia"
    return "Calibri"


def text_to_word(text_data: str | bytes, font: str = "Calibri", font_size: float = 12.0) -> bytes:
    """Convert raw or typed text into a beautifully formatted Microsoft Word (.docx) document."""
    if isinstance(text_data, bytes):
        try:
            raw_text = text_data.decode("utf-8")
        except UnicodeDecodeError:
            raw_text = text_data.decode("latin-1", errors="replace")
    else:
        raw_text = text_data

    target_font = _normalize_font_name(font)
    doc = Document()

    # Configure 1-inch margins
    for s in doc.sections:
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # Set document default style font
    try:
        normal_style = doc.styles["Normal"]
        f = getattr(normal_style, "font", None)
        if f:
            f.name = target_font
            f.size = Pt(font_size)
    except Exception:
        pass

    unicode_font_name = "Noto Sans Devanagari"

    for line in raw_text.splitlines():
        trimmed = line.rstrip()
        if not trimmed:
            # Add subtle empty line spacer
            p_empty = doc.add_paragraph()
            p_empty.paragraph_format.space_after = Pt(6)
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15

        r = p.add_run(trimmed)
        r.font.size = Pt(font_size)

        # Apply XML font properties for Latin + Complex Scripts
        is_unicode = has_devanagari(trimmed) or has_non_latin(trimmed)
        chosen_font = unicode_font_name if is_unicode else target_font

        r.font.name = chosen_font
        rPr = r._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = rPr.makeelement(qn("w:rFonts"), {})
            rPr.append(rFonts)

        rFonts.set(qn("w:ascii"), target_font)
        rFonts.set(qn("w:hAnsi"), target_font)
        rFonts.set(qn("w:cs"), unicode_font_name if is_unicode else target_font)

    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


def text_to_pdf(
    text_data: str | bytes,
    font: str = "Calibri",
    font_size: float = 12.0,
    searchable: bool = True,
) -> bytes:
    """Convert raw or typed text into a PDF, with searchable or flattened non-searchable output."""
    if isinstance(text_data, bytes):
        try:
            raw_text = text_data.decode("utf-8")
        except UnicodeDecodeError:
            raw_text = text_data.decode("latin-1", errors="replace")
    else:
        raw_text = text_data

    target_font = _normalize_font_name(font)
    reg_font = register_reportlab_font(target_font)
    unicode_font = register_reportlab_font("devanagari")

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        leftMargin=54,
        rightMargin=54,
        topMargin=54,
        bottomMargin=54,
    )

    story = []
    line_leading = font_size * 1.35

    for line in raw_text.splitlines():
        trimmed = line.rstrip()
        if not trimmed:
            story.append(Spacer(1, font_size * 0.75))
            continue

        escaped = html.escape(trimmed)
        use_font = unicode_font if (has_devanagari(trimmed) or has_non_latin(trimmed)) else reg_font

        ps = ParagraphStyle(
            name="CustomTextPara",
            fontName=use_font,
            fontSize=font_size,
            leading=line_leading,
            spaceAfter=4,
            wordWrap="CJK",
        )
        story.append(Paragraph(escaped, ps))

    if not story:
        story.append(Paragraph("&nbsp;", ParagraphStyle(name="Empty", fontName=reg_font, fontSize=font_size)))

    doc.build(story)
    pdf_bytes = buf.getvalue()

    if not searchable:
        pdf_bytes = make_non_searchable_pdf(pdf_bytes, dpi=300)

    return pdf_bytes

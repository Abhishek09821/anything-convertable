from __future__ import annotations
import base64, html, io, re
from reportlab.pdfgen import canvas as rl_canvas
from reportlab.lib.colors import HexColor
from reportlab.lib.utils import ImageReader
from docx import Document
from docx.shared import Pt, Inches
from PIL import Image as PILImage


# ── helpers ──────────────────────────────────────────────────────────────────

def _hex(color: str, fallback: str = '#000000') -> str:
    """Return a valid 6-digit hex string, defaulting on bad input."""
    s = (color or fallback).strip()
    if re.match(r'^#[0-9a-fA-F]{6}$', s):
        return s
    if re.match(r'^#[0-9a-fA-F]{3}$', s):
        r, g, b = s[1], s[2], s[3]
        return f'#{r}{r}{g}{g}{b}{b}'
    return fallback


def _data_url_to_bytes(src: str) -> bytes | None:
    """Convert a data:image/...;base64,... URL to raw bytes, or return None."""
    try:
        if src and src.startswith('data:'):
            _, encoded = src.split(',', 1)
            return base64.b64decode(encoded)
    except Exception:
        pass
    return None


def _image_reader(src: str):
    """Return a ReportLab ImageReader from a base64 data URL, or None."""
    raw = _data_url_to_bytes(src)
    if raw is None:
        return None
    try:
        return ImageReader(io.BytesIO(raw))
    except Exception:
        return None


# ── PDF ───────────────────────────────────────────────────────────────────────

def to_pdf(doc: dict) -> tuple[bytes, str]:
    buf = io.BytesIO()
    c = rl_canvas.Canvas(buf)
    for p in doc['pages']:
        w, h = float(p['width']), float(p['height'])
        c.setPageSize((w, h))
        # fill background
        bg = _hex(p.get('background', '#ffffff'), '#ffffff')
        c.setFillColor(HexColor(bg))
        c.rect(0, 0, w, h, fill=1, stroke=0)

        for e in sorted(p['elements'], key=lambda x: x.get('zIndex', 0)):
            b = e['bounds']
            ex, ey, ew, eh = float(b['x']), float(b['y']), float(b['width']), float(b['height'])
            # ReportLab y=0 is bottom; AEDOM y=0 is top → flip
            rl_y = h - ey - eh

            if e['type'] == 'image':
                reader = _image_reader(e.get('src', ''))
                if reader:
                    try:
                        c.drawImage(reader, ex, rl_y, ew, eh, preserveAspectRatio=False, mask='auto')
                        print(f'[export/pdf] drew image element id={e.get("id")} {ew:.0f}x{eh:.0f}')
                    except Exception as exc:
                        print(f'[export/pdf] image draw failed: {exc}')
                else:
                    # fallback: grey rect
                    c.setFillColor(HexColor('#e2e8f0'))
                    c.rect(ex, rl_y, ew, eh, fill=1, stroke=0)

            elif e['type'] == 'text':
                s = e.get('style', {})
                color = _hex(s.get('color', '#111111'))
                font_size = float(s.get('fontSize', 12))
                font_name = 'Helvetica-Bold' if s.get('fontWeight', 400) >= 700 else 'Helvetica'
                try:
                    c.setFillColor(HexColor(color))
                    c.setFont(font_name, font_size)
                    # draw at baseline (add fontSize to flip correctly)
                    c.drawString(ex, h - ey - font_size, e.get('text', ''))
                except Exception as exc:
                    print(f'[export/pdf] text draw failed: {exc}')

            elif e['type'] == 'shape':
                fill = _hex(e.get('fill', '#cccccc'))
                stroke = _hex(e.get('stroke', '#000000'))
                sw = float(e.get('strokeWidth', 0))
                c.setFillColor(HexColor(fill))
                c.setStrokeColor(HexColor(stroke))
                c.setLineWidth(sw)
                c.rect(ex, rl_y, ew, eh, fill=1, stroke=1 if sw > 0 else 0)

        c.showPage()
    c.save()
    print(f'[export/pdf] done, size={buf.tell()} bytes')
    return buf.getvalue(), 'application/pdf'


# ── DOCX ──────────────────────────────────────────────────────────────────────

def to_docx(doc: dict) -> tuple[bytes, str]:
    d = Document()
    for pi, p in enumerate(doc['pages']):
        if pi:
            d.add_page_break()
        els = sorted(p['elements'], key=lambda x: (x.get('zIndex', 0), x['bounds']['y'], x['bounds']['x']))
        for e in els:
            if e['type'] == 'image':
                raw = _data_url_to_bytes(e.get('src', ''))
                if raw:
                    try:
                        # Scale to fit within page width (6 inches)
                        img = PILImage.open(io.BytesIO(raw))
                        aspect = img.height / img.width if img.width else 1
                        w_in = min(6.0, e['bounds']['width'] / 96)
                        d.add_picture(io.BytesIO(raw), width=Inches(w_in))
                    except Exception as exc:
                        print(f'[export/docx] image failed: {exc}')
            elif e['type'] == 'text':
                s = e.get('style', {})
                para = d.add_paragraph()
                run = para.add_run(e.get('text', ''))
                run.bold = s.get('fontWeight', 400) >= 700
                run.italic = s.get('fontStyle', 'normal') == 'italic'
                run.font.size = Pt(s.get('fontSize', 12))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue(), 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'


# ── HTML ──────────────────────────────────────────────────────────────────────

def to_html(doc: dict) -> tuple[bytes, str]:
    pages = []
    for p in doc['pages']:
        els_html = []
        els = sorted(p['elements'], key=lambda x: x.get('zIndex', 0))
        for e in els:
            b = e['bounds']
            st = (
                f"position:absolute;"
                f"left:{b['x']}px;top:{b['y']}px;"
                f"width:{b['width']}px;height:{b['height']}px;"
                f"z-index:{e.get('zIndex', 0)};"
            )
            if e['type'] == 'image':
                src = e.get('src', '')
                if src:
                    els_html.append(
                        f'<img style="{st}object-fit:fill;" src="{src}" alt="{html.escape(e.get("alt",""))}" />'
                    )
            elif e['type'] == 'text':
                s = e.get('style', {})
                st += (
                    f"font-family:{html.escape(s.get('fontFamily','Arial'))};"
                    f"font-size:{s.get('fontSize',12)}px;"
                    f"font-weight:{s.get('fontWeight',400)};"
                    f"font-style:{s.get('fontStyle','normal')};"
                    f"color:{s.get('color','#111827')};"
                    f"overflow:hidden;"
                    f"white-space:pre-wrap;"
                )
                els_html.append(f'<div style="{st}">{html.escape(e.get("text",""))}</div>')
            elif e['type'] == 'shape':
                st += (
                    f"background:{e.get('fill','#cccccc')};"
                    f"border:{e.get('strokeWidth',0)}px solid {e.get('stroke','#000000')};"
                    f"border-radius:{12 if e.get('shape')=='roundRect' else 0}px;"
                )
                els_html.append(f'<div style="{st}"></div>')

        pages.append(
            f'<section style="position:relative;width:{p["width"]}px;height:{p["height"]}px;'
            f'background:{p.get("background","#fff")};margin:24px auto;overflow:hidden;">'
            + ''.join(els_html)
            + '</section>'
        )

    body = ''.join(pages)
    markup = (
        '<!doctype html><html><head><meta charset="utf-8">'
        '<title>Anything Convertable Export</title></head>'
        f'<body style="margin:0;background:#e5e7eb">{body}</body></html>'
    )
    return markup.encode('utf-8'), 'text/html'


# ── SVG ───────────────────────────────────────────────────────────────────────

def to_svg(doc: dict) -> tuple[bytes, str]:
    p = doc['pages'][0]
    w, h = p['width'], p['height']
    items = []
    els = sorted(p['elements'], key=lambda x: x.get('zIndex', 0))
    for e in els:
        b = e['bounds']
        ex, ey, ew, eh = b['x'], b['y'], b['width'], b['height']
        if e['type'] == 'image':
            src = e.get('src', '')
            if src:
                items.append(
                    f'<image x="{ex}" y="{ey}" width="{ew}" height="{eh}" '
                    f'preserveAspectRatio="none" href="{src}"/>'
                )
        elif e['type'] == 'text':
            s = e.get('style', {})
            fw = s.get('fontWeight', 400)
            fi = s.get('fontStyle', 'normal')
            items.append(
                f'<text x="{ex}" y="{ey + s.get("fontSize", 12)}" '
                f'font-family="{html.escape(s.get("fontFamily","Arial"))}" '
                f'font-size="{s.get("fontSize",12)}" '
                f'font-weight="{fw}" '
                f'font-style="{fi}" '
                f'fill="{s.get("color","#111827")}">'
                f'{html.escape(e.get("text",""))}</text>'
            )
        elif e['type'] == 'shape':
            items.append(
                f'<rect x="{ex}" y="{ey}" width="{ew}" height="{eh}" '
                f'fill="{e.get("fill","#cccccc")}" '
                f'stroke="{e.get("stroke","#000000")}" '
                f'stroke-width="{e.get("strokeWidth",0)}"/>'
            )

    svg = (
        f'<svg xmlns="http://www.w3.org/2000/svg" '
        f'xmlns:xlink="http://www.w3.org/1999/xlink" '
        f'width="{w}" height="{h}" viewBox="0 0 {w} {h}">'
        + ''.join(items)
        + '</svg>'
    )
    return svg.encode('utf-8'), 'image/svg+xml'

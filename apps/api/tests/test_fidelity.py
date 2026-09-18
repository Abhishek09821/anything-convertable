"""Pixel and content regressions for the accuracy-focused conversion pipeline."""
import io
import json
import fitz
import pytest
from PIL import Image, ImageDraw
from docx import Document
from pptx import Presentation


def image_bytes(image, fmt="PNG", **kwargs):
    out = io.BytesIO(); image.save(out, format=fmt, **kwargs); return out.getvalue()


def test_png_pixels_are_lossless():
    from app.conversions.image_to_pdf import convert
    image = Image.new("RGB", (80, 80), "white"); draw = ImageDraw.Draw(image)
    for n in range(0, 80, 2):
        draw.line((n, 0, n, 79), fill=(n * 3, 0, 255))
    with fitz.open(stream=convert(image_bytes(image)), filetype="pdf") as pdf:
        extracted = Image.open(io.BytesIO(pdf.extract_image(pdf[0].get_images()[0][0])["image"]))
        assert extracted.convert("RGB").tobytes() == image.tobytes()


def test_jpeg_stream_not_recompressed():
    from app.conversions.image_to_pdf import convert
    source = image_bytes(Image.new("RGB", (80, 80), "red"), "JPEG")
    with fitz.open(stream=convert(source), filetype="pdf") as pdf:
        assert pdf.extract_image(pdf[0].get_images()[0][0])["image"] == source


def test_palette_transparency_renders_white():
    from app.conversions.image_to_pdf import convert
    image = Image.new("P", (20, 20), 0)
    image.putpalette([0, 0, 0, 255, 0, 0] + [0] * 762); image.info["transparency"] = 0
    with fitz.open(stream=convert(image_bytes(image)), filetype="pdf") as pdf:
        assert pdf[0].get_pixmap().pixel(5, 5) == (255, 255, 255)


def test_all_tiff_pages_survive():
    from app.conversions.image_to_pdf import convert
    one = Image.new("RGB", (100, 80), "red"); two = Image.new("RGB", (50, 120), "blue")
    data = image_bytes(one, "TIFF", save_all=True, append_images=[two])
    with fitz.open(stream=convert(data), filetype="pdf") as pdf:
        assert len(pdf) == 2
        assert pdf[0].rect.width / pdf[0].rect.height == pytest.approx(1.25)
        assert pdf[1].rect.width / pdf[1].rect.height == pytest.approx(50 / 120)


def test_exif_rotates_dimensions_and_density():
    from app.conversions.image_to_pdf import convert
    exif = Image.Exif(); exif[274] = 6
    data = image_bytes(Image.new("RGB", (300, 100), "red"), "JPEG", exif=exif, dpi=(300, 100))
    with fitz.open(stream=convert(data), filetype="pdf") as pdf:
        assert pdf[0].rect.width == pytest.approx(72)
        assert pdf[0].rect.height == pytest.approx(72)


def test_pdf_slides_preserve_complete_graphics_and_sizes():
    from app.conversions.pdf_to_ppt import convert
    with fitz.open() as pdf:
        p = pdf.new_page(width=300, height=200)
        p.draw_rect((20, 20, 120, 100), fill=(1, 0, 0)); p.insert_text((30, 150), "Original text", fontsize=17)
        pdf.new_page(width=200, height=300).draw_circle((80, 90), 40, fill=(0, 0, 1))
        deck = Presentation(io.BytesIO(convert(pdf.tobytes(), fidelity="appearance")))
        assert deck.slide_width / deck.slide_height == pytest.approx(1.5)
        for index, slide in enumerate(deck.slides):
            picture = slide.shapes[0]
            assert picture.width / picture.height == pytest.approx(pdf[index].rect.width / pdf[index].rect.height)
            expected = pdf[index].get_pixmap(dpi=300, colorspace=fitz.csRGB, alpha=False)
            assert Image.open(io.BytesIO(picture.image.blob)).tobytes() == expected.samples


def test_pdf_span_merge_respects_font_family():
    from app.conversions.pdf_to_ppt import _merge_line_spans
    spans = [dict(text="First", bbox=(0, 0, 20, 12), size=12, font="Arial"),
             dict(text="Second", bbox=(21, 0, 50, 12), size=12, font="Georgia")]
    assert len(_merge_line_spans(spans)) == 2


def test_mixed_native_and_scanned_pages_preserve_text_and_scan(monkeypatch):
    from app.conversions.pdf_to_word import convert
    from app.conversions import ocr
    monkeypatch.setattr(ocr, "scan_lines", lambda page: [{"text": "Recognized scan", "bbox": (20, 20, 140, 36)}])
    with fitz.open() as pdf:
        pdf.new_page(width=300, height=400).insert_text((30, 50), "Native first page")
        page = pdf.new_page(width=300, height=400)
        page.insert_image(page.rect, stream=image_bytes(Image.new("RGB", (300, 400), "red")))
        document = Document(io.BytesIO(convert(pdf.tobytes())))
    text = " ".join(p.text for p in document.paragraphs)
    assert "Native first page" in text and "Recognized scan" in text
    assert len(document.inline_shapes) >= 1


def test_non_hindi_scripts_not_assigned_hindi_font():
    from app.conversions.pdf_to_word import _apply_font_to_docx
    doc = Document(); run = doc.add_paragraph("Greek Ελληνικά and Arabic العربية").runs[0]; run.font.name = "Arial"
    out = io.BytesIO(); doc.save(out)
    result = Document(io.BytesIO(_apply_font_to_docx(out.getvalue(), "original")))
    assert result.paragraphs[0].runs[0].font.name == "Arial"


def test_office_override_reaches_headers_and_preserves_emphasis():
    from app.conversions.office import apply_office_font
    doc = Document(); doc.add_paragraph("Styled text").runs[0].bold = True
    doc.sections[0].header.paragraphs[0].text = "Header"
    out = io.BytesIO(); doc.save(out)
    result = Document(io.BytesIO(apply_office_font(out.getvalue(), "Georgia")))
    assert result.paragraphs[0].runs[0].font.name == "Georgia"
    assert result.paragraphs[0].runs[0].bold
    assert result.sections[0].header.paragraphs[0].runs[0].font.name == "Georgia"


def test_native_office_export_uses_lossless_settings(monkeypatch):
    from pathlib import Path
    from types import SimpleNamespace
    from app.conversions import office
    monkeypatch.setattr(office, "office_binary", lambda: "/fake/soffice")
    profiles = []
    def run(args, **kwargs):
        profiles.append(next(a for a in args if a.startswith("-env:")))
        options = json.loads(args[args.index("--convert-to")+1].split(":", 2)[2])
        assert options["UseLosslessCompression"]["value"] == "true"
        assert options["ReduceImageResolution"]["value"] == "false"
        with fitz.open() as pdf:
            pdf.new_page().insert_text((50, 50), "Native Office output")
            (Path(args[args.index("--outdir")+1]) / "document.pdf").write_bytes(pdf.tobytes())
        return SimpleNamespace(returncode=0)
    monkeypatch.setattr(office.subprocess, "run", run)
    for _ in range(2):
        with fitz.open(stream=office.render_office(b"test", "docx"), filetype="pdf") as pdf:
            assert "Native Office output" in pdf[0].get_text()
    assert profiles[0] != profiles[1]


def test_actual_ocr_recovers_scan_text():
    import shutil
    if not shutil.which("tesseract"): pytest.skip("Tesseract is not installed")
    from app.conversions.ocr import scan_lines
    with fitz.open() as pdf:
        page = pdf.new_page(width=400, height=300); page.insert_text((40, 80), "Invoice total 12345", fontsize=22)
        pix = page.get_pixmap(dpi=300)
        with fitz.open() as scan:
            p = scan.new_page(width=400, height=300); p.insert_image(p.rect, stream=pix.tobytes("png"))
            assert "12345" in " ".join(x["text"] for x in scan_lines(p))


@pytest.fixture
def client():
    from fastapi.testclient import TestClient
    from app.main import app
    return TestClient(app)


def test_detects_misnamed_png(client):
    response = client.post("/v1/detect", files={"file": ("wrong.pdf", image_bytes(Image.new("RGB", (20,20))), "application/pdf")})
    assert response.json()["ext"] == ".png"


def test_detects_office_content_not_extension(client):
    doc = Document(); doc.add_paragraph("Hello"); out = io.BytesIO(); doc.save(out)
    response = client.post("/v1/detect", files={"file": ("wrong.pptx", out.getvalue())})
    assert response.json()["ext"] == ".docx"


def test_rejects_wrong_converter(client):
    response = client.post("/v1/convert/pdf_to_word", files={"file": ("wrong.pdf", image_bytes(Image.new("RGB", (20,20))))})
    assert response.status_code == 422


def test_unicode_filename_and_quality_notes(client):
    with fitz.open() as pdf:
        pdf.new_page().insert_text((30,50), "Some text")
        response = client.post("/v1/convert/pdf_to_ppt?fidelity=appearance", files={"file": ("हिंदी.pdf", pdf.tobytes())})
    assert response.status_code == 200
    assert "filename*=UTF-8''" in response.headers["content-disposition"]
    assert "not editable" in response.headers["x-conversion-warnings"]


def test_native_word_keeps_headers_images_and_font_override():
    from app.conversions.office import office_binary
    if not office_binary(): pytest.skip("LibreOffice integration requires installation")
    from app.conversions.word_to_pdf import convert
    from app.conversions.quality import begin_report, get_notes
    from app.conversions.fonts import get_font_path
    begin_report()
    from docx.shared import Inches
    doc = Document(); doc.sections[0].header.paragraphs[0].text = "Header must survive"
    p = doc.add_paragraph("Before picture ")
    p.add_run().add_picture(io.BytesIO(image_bytes(Image.new("RGB", (100,60), "red"))), width=Inches(1))
    p.add_run(" after picture")
    doc.add_picture(io.BytesIO(image_bytes(Image.new("RGB", (80,40), "blue"))), width=Inches(1))
    out = io.BytesIO(); doc.save(out)
    with fitz.open(stream=convert(out.getvalue(), font="Georgia"), filetype="pdf") as pdf:
        text = " ".join(p.get_text() for p in pdf)
        assert "Header must survive" in text and "Before picture" in text
        # PDF extractors can insert spaces between kerned glyphs. Check all actual characters.
        assert "afterpicture" in "".join(text.split())
        assert sum(len(p.get_images()) for p in pdf) >= 2
        if get_font_path("Georgia"):
            assert any("Georgia" in f[3] for p in pdf for f in p.get_fonts())
        else:
            assert any("substituted Georgia" in note for note in get_notes())


def test_native_powerpoint_is_searchable_and_preserves_shape_color():
    from app.conversions.office import office_binary
    if not office_binary(): pytest.skip("LibreOffice integration requires installation")
    from app.conversions.ppt_to_pdf import convert
    from pptx.util import Inches
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.dml.color import RGBColor
    deck = Presentation(); slide = deck.slides.add_slide(deck.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(6), Inches(1)).text = "Searchable slide text"
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1), Inches(3), Inches(2), Inches(1))
    shape.fill.solid(); shape.fill.fore_color.rgb = RGBColor(255,0,0)
    out = io.BytesIO(); deck.save(out)
    with fitz.open(stream=convert(out.getvalue()), filetype="pdf") as pdf:
        assert "Searchable slide text" in pdf[0].get_text()
        assert pdf[0].get_pixmap(dpi=72).pixel(100,240) == (255,0,0)


def test_hidden_ocr_layer_is_not_mistaken_for_native_text(monkeypatch):
    from app.conversions.pdf_to_word import convert
    from app.conversions import ocr
    calls = []
    def recognize(page):
        calls.append(page.number)
        return [{"text": "Recovered visible scan", "bbox": (20,20,160,40)}]
    monkeypatch.setattr(ocr, "scan_lines", recognize)
    with fitz.open() as pdf:
        p = pdf.new_page(width=300, height=400)
        p.insert_image(p.rect, stream=image_bytes(Image.new("RGB", (300,400), "white")))
        p.insert_text((20,40), "Invisible OCR layer", render_mode=3)
        output = Document(io.BytesIO(convert(pdf.tobytes())))
    assert calls == [0]
    assert "Recovered visible scan" in " ".join(p.text for p in output.paragraphs)
